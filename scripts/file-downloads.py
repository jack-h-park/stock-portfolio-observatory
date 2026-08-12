"""File broker downloads from the inbox into the directory that owns them.

Drop a freshly downloaded file into `$STOCK_DATA_DIR/inbox/` and run this. It
opens the file, works out what it is FROM THE CONTENTS, gives it the name
`docs/data-sources.md` defines, and moves it. Anything it cannot identify stays
in the inbox and is named in the output.

WHY BY CONTENT. Every source filename in this project is typed by hand at
download time, and that has already gone wrong twice in ways nothing caught:
`Robinhood - Agentic - 20060716 Year-to-date.csv` sat beside the `20260716` it
was meant to be — harmless only because "newest" happens to sort 2026 above
2006 — and a Bithumb statement named `2025년1-7월` actually held
`2026-01-01~2026-07-31`. docs/data-sources.md already says filenames are never
trusted for what a document contains, and that the periods on disk today were
read out of the documents. This is that rule, automated, so it holds for the
next download rather than for the ones someone checked by hand.

The incoming name is therefore never read. Not for the broker, not for the
document type, and above all not for the period: 미래에셋 prints
`제공내역 2022/01/01 ~ 2023/12/31`, Toss prints `조회 기간`, Bithumb prints
`기간 :`, Fidelity prints `Date downloaded`, the Robinhood crypto statement
names its month, and a Gain/Loss report carries its account and as-of date in
the PDF title. Those are the periods that get written.

TWO EXPORTS DECLARE NO DATE AT ALL. Chase's CSVs carry per-row dates and an
`As of` column that trails the download by a day, and nothing that says when the
export was taken; the Bithumb .xlsx says which window it covers but not when it
was pulled, which is what decides whether that window is final. For those, and
only those, the file's own modification time is used — for a file that was just
downloaded into the inbox that IS the download time — and every line that
depends on it says so, so a stale copy dragged in from elsewhere is visible
rather than silent.

FOUR RULES, in the order they matter:

  Never guess.     An unidentified file stays put and the run says which signals
                   were missing. Filing something into the wrong directory under
                   a confident name is worse than leaving it alone, because the
                   next reader believes the name.
  Never overwrite. A destination that already exists is compared byte for byte:
                   identical means already filed, different is a conflict to
                   report and for a person to resolve.
  Idempotent.      A second run over the same inbox does nothing. Identity is
                   checked by content hash against the WHOLE destination
                   directory, not just the target name, so re-filing a document
                   that is already there under a different name is a skip rather
                   than a duplicate — which matters because the KR extractor
                   reads every `mirae-*` file it finds and would count a
                   duplicate twice.
  Say what it did. Every move prints `<from> → <to>` with the evidence the name
                   was built from.

Run with --dry-run first; it prints the same plan and moves nothing.
"""

import hashlib
import os
import re
import shutil
import sys
import unicodedata
from datetime import date, timedelta
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent


def load_local_env():
    """Mirror scripts/env.mjs: .env.local, then .env, and a real env var wins.

    The hourly push runs this from launchd, which is not a login shell and has
    read nothing. Without this the timed run would file into the built-in
    default directory while a hand-run used the configured one.
    """
    for name in (".env.local", ".env"):
        path = REPO_ROOT / name
        if not path.exists():
            continue
        for line in path.read_text(encoding="utf-8").lstrip("﻿").splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            match = re.match(r"^([A-Za-z_][A-Za-z0-9_]*)=(.*)$", line)
            if not match or os.environ.get(match.group(1)) is not None:
                continue
            os.environ[match.group(1)] = match.group(2).strip("'\"")


load_local_env()

DATA_DIR = Path(os.environ.get("STOCK_DATA_DIR", Path.cwd() / "private-data"))
INBOX_DIR = DATA_DIR / "inbox"
PDF_PASSWORD = os.environ.get("STOCK_PDF_PASSWORD", "")

# The six directories docs/data-sources.md names. Keep in step with
# scripts/push-sources.sh, which pushes exactly these.
DIR_KR = "kr-statements"
DIR_US_HOLDINGS = "us-holdings"
DIR_US_TRANSACTIONS = "us-transactions"
DIR_US_TAX = "us-tax-documents"
DIR_BITHUMB = "crypto-bithumb"
DIR_RH_CRYPTO = "crypto-robinhood"

# What lands in the inbox but is not a source. macOS writes .DS_Store into any
# folder a Finder window has opened; a browser writes .crdownload/.part while a
# download is still arriving, and filing a half-written PDF would put a
# truncated statement into the pipeline under a confident name.
IGNORED_NAMES = {".DS_Store", "README.md", "Icon\r"}
IGNORED_SUFFIXES = {".crdownload", ".part", ".download", ".tmp"}

INBOX_README = """# inbox

Drop a freshly downloaded broker file here and run:

    make file-downloads          # or: pnpm file:downloads
    make file-downloads-dry      # shows the plan, moves nothing

The script reads the file, works out what it is from the CONTENTS — never from
the name you or the broker gave it — renames it to the convention in
docs/data-sources.md, and moves it to the directory that owns it. The hourly
push (`scripts/push-sources.sh`) runs it too, so a file dropped here is filed
and on the refresh host within the hour with no step for a person.

Anything it cannot identify STAYS HERE and is named in the output, with the
signals it looked for. That is deliberate: filing a document into the wrong
directory under a confident name is worse than leaving it alone, because the
next reader believes the name. Nothing in this directory is ever pushed or read
by the pipeline, so a file left here is inert, not lost.

Known exception: a Robinhood *transactions* CSV never says which strategy
account it came from — the export has no account column and no header — so it is
reported rather than filed. Name it by hand:
`us-transactions/robinhood-transactions-<agentic|longterm|midterm>-<date>.csv`.
"""


def nfc(value):
    """Composed Hangul, control characters gone.

    Two separate reasons, both already paid for elsewhere in this repo:
    Bithumb's PDF generator emits U+0001 BETWEEN glyphs, so `거래내역확인서`
    arrives as `거래내역\\x01확인서` and a literal comparison silently finds
    nothing; and the same text can arrive composed or decomposed, which compares
    unequal for identical Hangul.
    """
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))
    return unicodedata.normalize("NFC", text)


def despace(value):
    """Text with every space removed.

    Two documents need it for opposite reasons. The 미래에셋 certificates letter-
    space their titles (`거래내역 증 명 서`, `잔 고 증 명 서`), and the Chase 1099
    comes back with no spaces at all (`CONSOLIDATED2025FORMS1099ANDDETAILS`).
    Comparing both sides despaced is immune to either.
    """
    return re.sub(r"\s+", "", nfc(value))


def digits(value):
    return re.sub(r"\D", "", str(value or ""))


def iso(day):
    return day.strftime("%Y-%m-%d")


def compact(day):
    return day.strftime("%Y%m%d")


# ---------------------------------------------------------------------------
# Periods
#
# The period is the load-bearing part of the name: it is what decides whether a
# new download supersedes an old file or sits beside it, so it is never a
# SUMMARY of the coverage — it IS the coverage, to the day.
# ---------------------------------------------------------------------------

def period_from_window(start, end, allow_asof=True):
    """The period name for a window the document declared.

    A complete calendar year is `2025` and a run of them `2022-2023`; a window
    that starts on 1 January and stops short of the year end is an as-of
    (`20260716`), because the next download of the same thing covers the same
    ground plus more and only the newest may be read; anything else is the
    explicit window it is.

    `allow_asof` is off for the crypto sources, whose specs in
    scripts/source-files.mjs accept only a year or an explicit range: their
    exports are chosen windows that coexist, never a growing year-to-date.
    """
    whole_start = start.month == 1 and start.day == 1
    whole_end = end.month == 12 and end.day == 31
    if whole_start and whole_end:
        return str(start.year) if start.year == end.year else f"{start.year}-{end.year}"
    if allow_asof and whole_start and start.year == end.year:
        return compact(end)
    return f"{compact(start)}-{compact(end)}"


def period_from_rows(row_dates, downloaded):
    """The period for an export that names no window, only rows and a download date.

    Chase and Fidelity let you pick "2025", "year to date" or any window at all,
    and then say nothing about which you picked. What survives in the file is
    the rows.

    A year that has already ENDED is named as that year: the export is complete,
    nothing will be added to it, and archives coexist.

    ANYTHING REACHING INTO THE CURRENT YEAR IS NAMED BY THE ROWS IT ACTUALLY
    HAS, as a window. It used to be named for the download date — an as-of,
    which claims "everything up to here" and lets `pick: 'latest'` drop whatever
    it supersedes. That claim is not in evidence: a year-to-date export and a
    three-week window are the same file shape, and nothing in either says which
    it is. On 2026-08-11 a 12-row window covering 07-15…08-05 was named
    `chase-transactions-20260811.csv`, superseded 146 rows of year-to-date, and
    took eight checks down with it — the loss surfaced only as unrelated-looking
    failures somewhere else.

    A window states only what was observed. It cannot supersede anything, so a
    short download can no longer silently replace a long one; what it can do is
    OVERLAP one, and `us_transaction_periods_do_not_overlap` fails on that using
    the row dates rather than these names. The cost is that a year-to-date
    re-download now sits beside the file it repeats instead of replacing it, and
    has to be retired by hand — which that check names outright.
    """
    if not row_dates:
        return compact(downloaded)
    low, high = min(row_dates), max(row_dates)
    if high.year < downloaded.year:
        return str(low.year) if low.year == high.year else f"{low.year}-{high.year}"
    return f"{compact(low)}-{compact(high)}"


def parse_ymd(text, sep=r"[-/.]"):
    match = re.search(rf"(\d{{4}}){sep}(\d{{1,2}}){sep}(\d{{1,2}})", text)
    if not match:
        return None
    return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))


def parse_mdy(text):
    match = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", text)
    if not match:
        return None
    return date(int(match.group(3)), int(match.group(1)), int(match.group(2)))


# ---------------------------------------------------------------------------
# The document under the hand
# ---------------------------------------------------------------------------

class Document:
    """One inbox file, opened at most once per representation.

    Everything is lazy because most files answer to the first detector that
    looks at them, and a 619-page Toss statement is not worth extracting twice.
    """

    def __init__(self, path):
        self.path = path
        self.suffix = path.suffix.lower()
        self._pdf_pages = {}
        self._pdf_meta = None
        self._pdf_error = None
        self._lines = None
        self._cells = None
        self._digest = None
        self.encrypted = False

    # -- generic ----------------------------------------------------------
    @property
    def downloaded(self):
        """When this file arrived, as a date.

        Used only where the export itself declares nothing (see the module
        docstring). Every name built on it says so in the output.
        """
        return date.fromtimestamp(self.path.stat().st_mtime)

    @property
    def digest(self):
        if self._digest is None:
            self._digest = sha256_of(self.path)
        return self._digest

    # -- pdf --------------------------------------------------------------
    def _open_pdf(self):
        import logging

        import pdfplumber

        # pdfminer logs a FontBBox warning for every page of the Robinhood
        # statements. Sixty lines of it would bury the one thing this run has to
        # say, which is what moved where.
        logging.getLogger("pdfminer").setLevel(logging.ERROR)

        # The 미래에셋 종합 certificates and the 삼성 주식보상 listing are
        # encrypted; the ISA ones are not. Try open first so an unencrypted file
        # never depends on the password being configured, then fall back — the
        # same order scripts/extract-kr-statements.py uses.
        try:
            return pdfplumber.open(self.path)
        except Exception:
            if not PDF_PASSWORD:
                raise
            pdf = pdfplumber.open(self.path, password=PDF_PASSWORD)
            self.encrypted = True
            return pdf

    def page_text(self, index=0):
        """NFC text of one page, or '' if the file cannot be read as a PDF."""
        if self.suffix != ".pdf":
            return ""
        if index in self._pdf_pages:
            return self._pdf_pages[index]
        text = ""
        try:
            with self._open_pdf() as pdf:
                self._pdf_meta = pdf.metadata or {}
                if index < len(pdf.pages):
                    text = nfc(pdf.pages[index].extract_text() or "")
        except Exception as exc:  # unreadable, wrong password, not a PDF
            self._pdf_error = f"{type(exc).__name__}: {exc}"
            if not PDF_PASSWORD:
                # The 미래에셋 종합 certificates and the 삼성 listing are
                # encrypted, so an unset password looks exactly like a corrupt
                # file. Say which it probably is.
                self._pdf_error += " (STOCK_PDF_PASSWORD is not set — the 미래에셋 종합 and 삼성 certificates are encrypted)"
        self._pdf_pages[index] = text
        return text

    @property
    def pdf_meta(self):
        if self._pdf_meta is None:
            self.page_text(0)
        return self._pdf_meta or {}

    @property
    def pdf_error(self):
        self.page_text(0)
        return self._pdf_error

    # -- text/csv ---------------------------------------------------------
    @property
    def lines(self):
        if self._lines is None:
            if self.suffix in (".csv", ".txt"):
                raw = self.path.read_text(encoding="utf-8-sig", errors="replace")
                self._lines = raw.splitlines()
            else:
                self._lines = []
        return self._lines

    # -- xlsx -------------------------------------------------------------
    @property
    def cells(self):
        """The first sheet's top-left block, as {(row, col): text}."""
        if self._cells is None:
            self._cells = {}
            if self.suffix == ".xlsx":
                try:
                    import warnings

                    import openpyxl

                    with warnings.catch_warnings():
                        warnings.simplefilter("ignore")
                        book = openpyxl.load_workbook(self.path, read_only=True)
                    sheet = book.worksheets[0]
                    for r, row in enumerate(sheet.iter_rows(max_row=6, values_only=True), start=1):
                        for c, value in enumerate(row, start=1):
                            if value is not None:
                                self._cells[(r, c)] = nfc(value)
                    book.close()
                except Exception:
                    pass
        return self._cells

    def cell(self, row, col=1):
        return self.cells.get((row, col), "")

    # -- summary for the unidentified report ------------------------------
    def first_line(self):
        if self.suffix == ".pdf":
            text = self.page_text(0)
            if not text:
                return f"(no text: {self.pdf_error or 'empty first page'})"
            return text.split("\n")[0][:110]
        if self.cells:
            return self.cell(1, 1)[:110]
        for line in self.lines:
            if line.strip():
                return line[:110]
        return "(empty)"


def sha256_of(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


class Plan:
    """Where one inbox file goes, and the evidence that says so."""

    def __init__(self, subdir, name, evidence, group=None, order=None):
        self.subdir = subdir
        self.name = name
        self.evidence = evidence
        # Set by the families that arrive split across several files (Toss).
        # Members of one group share a destination name until the run has seen
        # all of them and can number them.
        self.group = group
        self.order = order


class Refusal:
    """Recognised, and deliberately not filed.

    Different from "no detector claimed it": the document is known, and what is
    missing is a fact the document does not contain. Saying which fact is the
    whole value — a bare "unidentified" would send the reader back to open the
    file and find the same nothing.
    """

    def __init__(self, what, reason, remedy=None):
        self.what = what
        self.reason = reason
        self.remedy = remedy


# ---------------------------------------------------------------------------
# Detectors. Each takes a Document and returns a Plan, a Refusal, or None
# ("not mine"). Order does not matter: the run fails loudly if two claim the
# same file, because that means a marker is not as specific as it looks.
# ---------------------------------------------------------------------------

# 미래에셋 계좌번호 → the account type its certificates belong to. Read off the
# certificates on disk rather than inferred from the number's shape: a 잔고증명서
# does not print 계좌유형, and guessing which account a balance belongs to would
# put ISA holdings under the 종합 account. An unknown account number is refused.
MIRAE_ACCOUNTS = {
    "456121492000": "isa",
    "220224249320": "general",
}


def mirae_period_window(text):
    """The 제공내역 window: `2022/01/01 ~ 2023/12/31`."""
    match = re.search(
        r"(\d{4})/(\d{2})/(\d{2})\s*~\s*(\d{4})/(\d{2})/(\d{2})", nfc(text)
    )
    if not match:
        return None
    g = [int(x) for x in match.groups()]
    return date(g[0], g[1], g[2]), date(g[3], g[4], g[5])


def detect_mirae_transactions(doc):
    """미래에셋 거래내역증명서 → kr-statements/mirae-<isa|general>-transactions-…

    The title is letter-spaced on the cover (`거래내역 증 명 서`), so it is
    matched despaced. 계좌유형 is not on the cover — it is in the table header
    that repeats on every data page — so page 2 supplies whether this is the ISA
    or the 종합 account.
    """
    cover = doc.page_text(0)
    if "거래내역증명서" not in despace(cover) or "미래에셋증권" not in despace(cover):
        return None

    window = mirae_period_window(cover)
    if not window:
        return Refusal(
            "미래에셋 거래내역증명서",
            "the cover has no `제공내역 YYYY/MM/DD ~ YYYY/MM/DD` window, and the "
            "period must come from the document",
        )
    start, end = window

    body = despace(doc.page_text(1))
    kind = None
    match = re.search(r"계좌유형(.*?)고객명", body)
    if match:
        kind = "isa" if "ISA" in match.group(1) else "general" if "종합" in match.group(1) else None
    if kind is None:
        account = re.search(r"계좌번호(\d[\d-]+)", body)
        kind = MIRAE_ACCOUNTS.get(digits(account.group(1))) if account else None
    if kind is None:
        return Refusal(
            "미래에셋 거래내역증명서",
            "page 2 names neither 계좌유형 ISA/종합 nor a 계좌번호 this repo has "
            "seen, so which account it belongs to is unknown",
            "add the account number to MIRAE_ACCOUNTS in scripts/file-downloads.py",
        )

    evidence = [f"제공내역 {iso(start)} ~ {iso(end)}", f"계좌유형 {kind}"]
    period = period_from_window(start, end)
    name = f"mirae-{kind}-transactions-{period}"

    # The 종합 certificates are issued in numbered batches — three of them came
    # out on one day covering adjacent windows — so the 발급번호 goes on the name
    # to keep two certificates for one window apart. The ISA ones arrive one per
    # year and have never needed it.
    if kind == "general":
        issue = re.search(r"NO\.(\d{4})-(\d{3})-(\d{8})", despace(cover))
        if not issue:
            return Refusal(
                "미래에셋 종합 거래내역증명서",
                "the cover has no `N O. YYYY-NNN-NNNNNNNN` 발급번호, which is what "
                "keeps two certificates for the same window apart",
            )
        serial = issue.group(3).lstrip("0")[-4:]
        evidence.append(f"발급번호 …{serial}")
        name = f"{name}-{serial}"

    return Plan(DIR_KR, f"{name}.pdf", evidence)


def detect_mirae_balance(doc):
    """미래에셋 잔고증명서 → kr-statements/mirae-<kind>-balance-<기준일자>-<id>.pdf

    A balance certificate is a position AT A MOMENT, so its period is the
    기준일자 it prints — not the 발급일시, which is only when someone asked for it
    and can be most of a year later (2025-09-11 balances issued 2026-07-16).
    """
    cover = despace(doc.page_text(0))
    if "잔고증명서" not in cover:
        return None

    as_of = None
    match = re.search(r"기준일자발급일시.*?(\d{4}-\d{2}-\d{2})", cover, re.S)
    if match:
        as_of = parse_ymd(match.group(1))
    if not as_of:
        return Refusal(
            "미래에셋 잔고증명서",
            "no 기준일자 on the cover, and a balance certificate's period is the "
            "date it is a balance AS OF",
        )

    account = re.search(r"계좌번호계좌명부기명실명확인번호(\d[\d-]+)", cover)
    kind = MIRAE_ACCOUNTS.get(digits(account.group(1))) if account else None
    if kind is None:
        return Refusal(
            "미래에셋 잔고증명서",
            "the 계좌번호 is not one this repo has seen, and a 잔고증명서 does not "
            "print 계좌유형 — which account it belongs to is unknown",
            "add the account number to MIRAE_ACCOUNTS in scripts/file-downloads.py",
        )

    issue = re.search(r"발급번호:(\d{4})-(\d{3})-(\d{8})", cover)
    if not issue:
        return Refusal(
            "미래에셋 잔고증명서",
            "no 발급번호 on the cover; two certificates can share a 기준일자 and "
            "that number is what tells them apart",
        )
    serial = issue.group(3).lstrip("0")[-4:]

    return Plan(
        DIR_KR,
        f"mirae-{kind}-balance-{compact(as_of)}-{serial}.pdf",
        [f"기준일자 {iso(as_of)}", f"발급번호 …{serial}", f"계좌 {kind}"],
    )


def detect_toss_transactions(doc):
    """토스증권 거래내역서 → kr-statements/toss-transactions-<period>[-NofM].pdf

    A long statement comes out of Toss split across several PDFs that share one
    발급번호 and one 조회 기간 and differ only in which rows they carry, so the
    part numbers are assigned across the whole run (see number_split_families)
    by first transaction date, never taken from the incoming names.
    """
    cover = doc.page_text(0)
    flat = despace(cover)
    if "거래내역서" not in flat or "발급번호" not in flat:
        return None
    account = re.search(r"계좌번호(\d{3})-", flat)
    if not account or account.group(1) != "137":
        return None

    match = re.search(
        r"조회기간(\d{4})년(\d{1,2})월(\d{1,2})일~(\d{4})년(\d{1,2})월(\d{1,2})일", flat
    )
    if not match:
        return Refusal(
            "토스증권 거래내역서",
            "no `조회 기간 YYYY년 M월 D일 ~ …` on the cover, and the period must "
            "come from the document",
        )
    g = [int(x) for x in match.groups()]
    start, end = date(g[0], g[1], g[2]), date(g[3], g[4], g[5])

    issue = re.search(r"발급번호([\w-]+)", flat)
    issue_no = issue.group(1) if issue else ""

    # Which slice of the statement this file is. The rows are in date order, so
    # the first transaction on page 1 orders the parts among themselves.
    first_row = None
    for line in cover.split("\n"):
        found = re.match(r"^(\d{4})\.(\d{2})\.(\d{2})\s", line)
        if found:
            first_row = date(*(int(x) for x in found.groups()))
            break

    period = period_from_window(start, end)
    return Plan(
        DIR_KR,
        f"toss-transactions-{period}.pdf",
        [f"조회 기간 {iso(start)} ~ {iso(end)}", f"발급번호 {issue_no}"],
        group=("toss", period, issue_no),
        order=(first_row or date(1900, 1, 1), doc.path.name),
    )


def detect_samsung_rsu(doc):
    """삼성증권 계좌거래내역 (주식보상) → kr-statements/samsung-rsu-transactions-<id>.pdf

    One listing per account rather than per period — it is the whole life of the
    주식보상 account — so the account, not a window, is what the name carries.
    """
    cover = despace(doc.page_text(0))
    if "계좌거래내역" not in cover or "주식보상" not in cover:
        return None
    account = re.search(r"계좌번호(\d[\d-]*)", cover)
    if not account:
        return Refusal(
            "삼성증권 계좌거래내역 (주식보상)",
            "no 계좌번호 on the first page, and the account is what this name "
            "carries in place of a period",
        )
    tail = digits(account.group(1))[-5:]
    window = re.search(r"조회일자(\d{4}-\d{2}-\d{2})~(\d{4}-\d{2}-\d{2})", cover)
    evidence = [f"계좌번호 …{tail}"]
    if window:
        evidence.append(f"조회일자 {window.group(1)} ~ {window.group(2)}")
    return Plan(DIR_KR, f"samsung-rsu-transactions-{tail}.pdf", evidence)


def detect_robinhood_crypto_statement(doc):
    """Robinhood Crypto Monthly Statement → crypto-robinhood/…-<YYYYMM>.pdf"""
    text = doc.page_text(0)
    head = text.split("\n")
    if not head or head[0].strip() != "Crypto Statement":
        return None
    match = re.search(r"^(\d{2})-(\d{4})$", "\n".join(head[:6]), re.M)
    if not match:
        return Refusal(
            "Robinhood Crypto statement",
            "the statement does not print its `MM-YYYY` month under the title",
        )
    month, year = match.group(1), match.group(2)
    return Plan(DIR_RH_CRYPTO, f"robinhood-crypto-statement-{year}{month}.pdf",
                [f"statement month {year}-{month}"])


def detect_robinhood_gain_loss(doc):
    """Robinhood Gain/Loss report → us-holdings/robinhood-holdings-<acct>-<asof>.pdf

    The account and the as-of date are in the PDF title
    (`1478 Unrealized as of 7/16/2026 created 7/16/2026`), which is metadata the
    generator wrote — not the filename someone typed. The page body carries no
    account number at all, which is why the title is what is read.

    The MCP snapshot has since taken over as the source of Robinhood lots, and
    these still belong here: they are the only record of those positions before
    the snapshot existed, `us_pdf_evidence_extracted` asserts on them, and
    scripts/extract-us-pdf-evidence.py still builds a lot's account label from
    the four digits in the name. They are also a download by the test in
    docs/data-sources.md — a human had to email customer support and wait for
    them — so `us-holdings/` is the right side of that boundary.
    """
    text = doc.page_text(0)
    if "OPEN LONGS" not in text or "WS Cost Adj" not in text:
        return None
    title = str(doc.pdf_meta.get("Title") or "")
    match = re.match(r"\s*(\d{4})\s+Unrealized as of\s+(\d{1,2})/(\d{1,2})/(\d{4})", title)
    if not match:
        return Refusal(
            "Robinhood Gain/Loss report",
            f"the PDF title ({title!r}) does not carry the "
            "`<account> Unrealized as of M/D/YYYY` the account and date are read from",
        )
    account = match.group(1)
    as_of = date(int(match.group(4)), int(match.group(2)), int(match.group(3)))
    return Plan(
        DIR_US_HOLDINGS,
        f"robinhood-holdings-{account}-{compact(as_of)}.pdf",
        [f"pdf title: account {account}, unrealized as of {iso(as_of)}"],
    )


# 1099 packages. Each broker issues one per year; Fidelity issues one PER
# ACCOUNT and two arrive every year, so only its name carries the account.
TAX_FORMS = [
    {
        "broker": "fidelity",
        "label": "Fidelity consolidated 1099",
        "marks": ("TAXREPORTINGSTATEMENT", "FIDELITYBROKERAGESERVICES"),
        "year": r"(\d{4})TAXREPORTINGSTATEMENT",
        # `Account No. Z37-480490` — the last four digits are what the name
        # carries, and the branch prefix differs between the two accounts.
        "account": r"AccountNo\.([A-Z0-9-]+)",
    },
    {
        "broker": "chase",
        "label": "J.P. Morgan consolidated 1099",
        "marks": ("J.P.MORGANSECURITIESLLC", "FORMS1099"),
        "year": r"CONSOLIDATED(\d{4})FORMS1099",
        "account": None,
    },
    {
        "broker": "robinhood",
        "label": "Robinhood consolidated 1099",
        "marks": ("RobinhoodMarketsInc", "1099"),
        # 2024 says "Consolidated Tax Statement", 2025 dropped the word.
        "year": r"Enclosedisyour(\d{4})(?:Consolidated)?TaxStatement",
        "account": None,
    },
]


def detect_tax_form(doc):
    """A consolidated 1099 → us-tax-documents/<broker>-1099-[<account>-]<year>.pdf"""
    flat = despace(doc.page_text(0))
    if "1099" not in flat:
        return None
    for form in TAX_FORMS:
        if not all(mark in flat for mark in form["marks"]):
            continue
        year = re.search(form["year"], flat)
        if not year:
            return Refusal(
                form["label"],
                "the cover page does not print the tax year in the shape this "
                f"reads ({form['year']})",
            )
        parts = [form["broker"], "1099"]
        evidence = [f"tax year {year.group(1)}"]
        if form["account"]:
            account = re.search(form["account"], flat)
            if not account:
                return Refusal(
                    form["label"],
                    "no account number on the cover; Fidelity issues one 1099 per "
                    "account and two arrive each year, so the name needs it",
                )
            tail = digits(account.group(1))[-4:]
            if len(tail) < 4:
                return Refusal(
                    form["label"],
                    f"the account on the cover ({account.group(1)!r}) has fewer "
                    "than four digits to name it by",
                )
            parts.append(tail)
            evidence.append(f"account …{tail}")
        parts.append(year.group(1))
        return Plan(DIR_US_TAX, "-".join(parts) + ".pdf", evidence)
    return None


def detect_bithumb_statement(doc):
    """빗썸 거래내역확인서 → crypto-bithumb/bithumb-statement-<period>[-partial].pdf"""
    text = doc.page_text(0)
    flat = despace(text)
    if "거래내역확인서" not in flat:
        return None
    window = re.search(r"(\d{4}-\d{2}-\d{2})~(\d{4}-\d{2}-\d{2})", flat)
    if not window:
        return Refusal(
            "빗썸 거래내역확인서",
            "no `조회기간 YYYY-MM-DD~YYYY-MM-DD` on the first page, and the period "
            "must come from the document — a Bithumb filename has already lied "
            "about exactly this",
        )
    start, end = parse_ymd(window.group(1)), parse_ymd(window.group(2))
    evidence = [f"조회기간 {iso(start)} ~ {iso(end)}"]

    issued = None
    stamp = re.search(r"발급일자:(\d{4}-\d{2}-\d{2})", flat)
    if stamp:
        issued = parse_ymd(stamp.group(1))
        evidence.append(f"발급일자 {iso(issued)}")
    name = f"bithumb-statement-{period_from_window(start, end, allow_asof=False)}"
    if issued and end >= issued:
        # The window runs to the day it was issued or past it, so it cannot be
        # final — the next download extends it. Same fact 빗썸's own 일부 marks.
        name += "-partial"
        evidence.append("window reaches the issue date → not final")
    return Plan(DIR_BITHUMB, f"{name}.pdf", evidence)


def detect_bithumb_activity(doc):
    """빗썸 기간별 거래 내역 (.xlsx) → crypto-bithumb/bithumb-activity-<period>[-partial].xlsx"""
    if doc.suffix != ".xlsx" or "Bithumb" not in doc.cell(1, 1):
        return None
    header = doc.cell(2, 1)
    window = re.search(
        r"기간\s*:\s*(\d{4}-\d{2}-\d{2})\s*\d{2}:\d{2}:\d{2}\s*~\s*(\d{4}-\d{2}-\d{2})\s*(\d{2}:\d{2}:\d{2})",
        nfc(header),
    )
    if not window:
        return Refusal(
            "빗썸 기간별 거래 내역",
            f"row 2 is {header!r}, not the `기간 : <from> ~ <to>` the period is read from",
        )
    start, end = parse_ymd(window.group(1)), parse_ymd(window.group(2))

    # Bithumb renders a window's end as the START of the following one — a
    # Sep–Dec export ends `2025-01-01 00:00:59`. So an end that lands on the
    # first of a month is exclusive and the last covered day is the day before;
    # an end anywhere else is the day the export was actually asked to stop.
    evidence = [f"기간 {iso(start)} ~ {iso(end)} {window.group(3)}"]
    if end.day == 1:
        end -= timedelta(days=1)
        evidence.append(f"end is a month boundary (exclusive) → covers to {iso(end)}")

    name = f"bithumb-activity-{period_from_window(start, end, allow_asof=False)}"
    taken = doc.downloaded
    if end >= taken:
        # The .xlsx, unlike the PDF, prints no issue date, so when it was pulled
        # has to come from the file itself — see the module docstring.
        name += "-partial"
        evidence.append(
            f"window reaches the download ({iso(taken)}, from the file's timestamp) → not final"
        )
    return Plan(DIR_BITHUMB, f"{name}.xlsx", evidence)


def csv_head(doc, count=6):
    return "\n".join(doc.lines[:count])


def detect_chase(doc):
    """Chase CSVs → us-holdings/chase-holdings-… or us-transactions/chase-transactions-…

    Chase declares no export date anywhere: the holdings export's `As of` column
    is the previous close and the transactions export carries only row dates. So
    the as-of is the download — the file's own timestamp — and the output says
    so on the line.
    """
    if doc.suffix != ".csv":
        return None
    head = csv_head(doc, 3)
    taken = doc.downloaded
    stamp = f"download {iso(taken)} (from the file's timestamp; Chase prints no export date)"

    if head.startswith("Account name,Account number,Account type") and "As of" in head:
        as_of = None
        for line in doc.lines[1:3]:
            found = re.search(r'"(\d{2}/\d{2}/\d{4}) \d{2}:\d{2}:\d{2}"', line)
            if found:
                as_of = parse_mdy(found.group(1))
                break
        evidence = [stamp] + ([f"priced as of {iso(as_of)}"] if as_of else [])
        return Plan(DIR_US_HOLDINGS, f"chase-holdings-{compact(taken)}.csv", evidence)

    if head.startswith("Trade Date,Post Date,Settlement Date") and "Cusip" in head:
        rows = []
        for line in doc.lines[1:]:
            found = re.match(r'^"(\d{1,2}/\d{1,2}/\d{4})"', line)
            if found:
                rows.append(parse_mdy(found.group(1)))
        period = period_from_rows(rows, taken)
        evidence = [stamp]
        if rows:
            evidence.append(f"rows {iso(min(rows))} … {iso(max(rows))}")
        return Plan(DIR_US_TRANSACTIONS, f"chase-transactions-{period}.csv", evidence)

    return None


def detect_fidelity(doc):
    """Fidelity transactions CSV → us-transactions/fidelity-transactions-<period>.csv

    Fidelity is the well-behaved one: it prints `Date downloaded 07/23/2026` in
    its own footer, so nothing here depends on the filesystem.
    """
    if doc.suffix != ".csv":
        return None
    if "Run Date,Action,Symbol,Description" not in csv_head(doc, 8):
        return None
    downloaded = None
    for line in reversed(doc.lines[-12:]):
        if "Date downloaded" in line:
            downloaded = parse_mdy(line)
            break
    if not downloaded:
        return Refusal(
            "Fidelity transactions export",
            "the footer has no `Date downloaded MM/DD/YYYY`, which is what says "
            "whether this is a complete year or a year-to-date snapshot",
        )
    rows = []
    for line in doc.lines:
        found = re.match(r"^(\d{1,2}/\d{1,2}/\d{4}),", line)
        if found:
            rows.append(parse_mdy(found.group(1)))
    period = period_from_rows(rows, downloaded)
    evidence = [f"Date downloaded {iso(downloaded)}"]
    if rows:
        evidence.append(f"rows {iso(min(rows))} … {iso(max(rows))}")
    return Plan(DIR_US_TRANSACTIONS, f"fidelity-transactions-{period}.csv", evidence)


# Merrill's site exports positions under two different layouts, and which one
# you get depends on a control on the page rather than on anything the file says
# about itself. Both are holdings:
#
#   tax-lot detail   "Symbol ","Quantity","Unit Cost","Cost Basis",…   plus a
#                    per-position `Acquisition Date` block underneath each row.
#   flat positions   "Symbol ","Description","Quantity","Price",…      with the
#                    basis under `Total Client Investment` and no lots at all.
#
# So `Cost Basis` is not the marker for "these are holdings" — it is the marker
# for ONE of the two layouts, and requiring it rejected a perfectly good
# positions export as ambiguous. What both layouts share is the `Symbol ` column
# (trailing space and all, which is Merrill's, not a typo here) and a
# basis column under one of two names. Naming the basis column in the evidence
# line is what says which layout arrived, since the destination name cannot.
# Merrill respells its own headers between exports — `Total client investment`
# for `Total Client Investment` — so these are compared case-insensitively.
MERRILL_BASIS_COLUMNS = ('"Cost Basis"', '"Total Client Investment"')


MONTH_ABBREVIATIONS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


def parse_mon_dd_yyyy(text):
    """`Jul-31-2026` → date(2026, 7, 31).

    Fidelity stamps its two exports with two different date formats: the
    transactions footer says `Date downloaded 07/31/2026 05:22 pm` and the
    positions footer says `Date downloaded Jul-31-2026 at 8:24 p.m ET`. Same
    broker, same sentence, same afternoon — so one parser cannot serve both, and
    reusing parse_mdy on the positions file would find no date and refuse a file
    that does in fact declare one.
    """
    match = re.search(r"([A-Za-z]{3})[a-z]*-(\d{1,2})-(\d{4})", text)
    if not match:
        return None
    month = MONTH_ABBREVIATIONS.get(match.group(1).lower())
    if not month:
        return None
    return date(int(match.group(3)), month, int(match.group(2)))


def detect_fidelity_positions(doc):
    """Fidelity positions CSV → us-holdings/fidelity-holdings-<asof>.csv

    A positions export is a snapshot at a moment, so its period is an as-of and
    only the newest may be read — the same rule the Chase and Merrill holdings
    exports follow.

    ONE ACCOUNT PER FILE, ENFORCED. Every row carries its own `Account number`,
    so a multi-account export parses perfectly well; what it cannot do is be
    NAMED. The holdings grammar has no account slot in use here, so two accounts
    filed under one `fidelity-holdings-<date>.csv` would be indistinguishable
    from one account's export of the same date — and `pick: 'latest'` would then
    let a later single-account download silently supersede the pair, dropping an
    account's positions with nothing to say so. That is this repo's recurring
    failure: not an error, just a smaller portfolio than exists. So a
    multi-account export is refused with the remedy, rather than filed under a
    name that understates it.
    """
    if doc.suffix != ".csv" or not doc.lines:
        return None
    # utf-8-sig on the read already ate the BOM this file carries; the header is
    # matched on the columns that make it a POSITIONS export rather than the
    # transactions one, which starts `Run Date,Action,…`.
    header = doc.lines[0]
    if not header.startswith("Account number,Account name,Symbol"):
        return None
    if "Cost basis total" not in header:
        return Refusal(
            "Fidelity positions export",
            "the header has no `Cost basis total` column, which is the only "
            "basis this export carries and the reason it is worth ingesting",
        )

    downloaded = None
    for line in reversed(doc.lines[-12:]):
        if "Date downloaded" in line:
            downloaded = parse_mon_dd_yyyy(line)
            break
    if not downloaded:
        return Refusal(
            "Fidelity positions export",
            "the footer has no `Date downloaded <Mon>-<DD>-<YYYY>`, and a "
            "positions snapshot's period is the moment it was taken",
        )

    accounts = []
    for line in doc.lines[1:]:
        found = re.match(r"^([A-Z0-9]{6,}),", line)
        if found and found.group(1) not in accounts:
            accounts.append(found.group(1))
    if not accounts:
        return Refusal(
            "Fidelity positions export",
            "no row carries an account number, so there is nothing to check the "
            "one-account-per-file rule against",
        )
    if len(accounts) > 1:
        return Refusal(
            "Fidelity positions export",
            f"it holds {len(accounts)} accounts ({', '.join(accounts)}) and the "
            "holdings name carries no account, so filing it would understate it",
            "re-export one account at a time, or split the CSV by its "
            "`Account number` column and drop the parts in separately",
        )

    return Plan(
        DIR_US_HOLDINGS,
        f"fidelity-holdings-{compact(downloaded)}.csv",
        [f"Date downloaded {iso(downloaded)}", f"account {accounts[0]}"],
    )


def detect_merrill(doc):
    """Merrill CSVs → us-holdings/merrill-holdings-… or us-transactions/merrill-transactions-…

    Both start `Exported on: 07/15/2026 08:53 PM ET`, which is the as-of, and
    differ in the table header further down.
    """
    if doc.suffix != ".csv" or not doc.lines:
        return None
    # The 2026-08-11 export quotes this line where the older ones left it bare.
    # Same sentence, same date, one character of CSV quoting between the file
    # being recognised and being dropped on the floor as unidentified.
    first = doc.lines[0].lstrip('"')
    if not first.startswith("Exported on:"):
        return None
    exported = parse_mdy(first)
    if not exported:
        return Refusal(
            "Merrill export",
            f"the first line is {doc.lines[0]!r} but carries no MM/DD/YYYY date",
        )
    head = csv_head(doc, 12)
    evidence = [f"Exported on {iso(exported)}"]
    if ('"Trade Date"' in head and '"Settlement Date"' in head) or (
        '"Settlement date"' in head and '"Description"' in head and '"Symbol/CUSIP"' in head
    ):
        return Plan(DIR_US_TRANSACTIONS, f"merrill-transactions-{compact(exported)}.csv", evidence)
    # `Symbol` in the older layouts, `Positions` in the 2026-08 one. The table
    # underneath is the same; only its heading was renamed.
    if '"Symbol ' in head or '"Positions"' in head:
        basis = next((c for c in MERRILL_BASIS_COLUMNS if c.lower() in head.lower()), None)
        if basis is None:
            return Refusal(
                "Merrill holdings export",
                "the `Symbol` column is there but neither basis column is "
                f"({' nor '.join(MERRILL_BASIS_COLUMNS)}), and without a basis "
                "this is not a holdings export this repo can read",
            )
        evidence.append(f"holdings layout: basis under {basis}")
        return Plan(DIR_US_HOLDINGS, f"merrill-holdings-{compact(exported)}.csv", evidence)
    return Refusal(
        "Merrill export",
        "neither the transactions header (`Trade Date`, `Settlement Date`) nor "
        "the holdings header (`Symbol`) is present, so which of the two this is "
        "cannot be told",
    )


def detect_robinhood_transactions(doc):
    """Robinhood transactions CSV — recognised, and deliberately not filed.

    The export is nine columns of activity and nothing else: no account number,
    no header, no footer, no metadata. Which of the three strategy accounts it
    came from exists only in the name the person typed, and that name is exactly
    what this script refuses to believe. Those tokens are half the `account` on
    every Robinhood row in the database, so filing the wrong one would mislabel
    real holdings — the one outcome worse than leaving the file alone.
    """
    if doc.suffix != ".csv" or not doc.lines:
        return None
    if not doc.lines[0].startswith('"Activity Date","Process Date","Settle Date","Instrument"'):
        return None
    rows = []
    for line in doc.lines[1:]:
        found = re.match(r'^"(\d{1,2}/\d{1,2}/\d{4})"', line)
        if found:
            rows.append(parse_mdy(found.group(1)))
    span = f", rows {iso(min(rows))} … {iso(max(rows))}" if rows else ""
    return Refusal(
        "Robinhood transactions export",
        f"the CSV never says which strategy account it came from{span} — no "
        "account column, no header, no footer",
        "name it by hand: us-transactions/robinhood-transactions-"
        "<agentic|longterm|midterm>-<period>.csv",
    )


DETECTORS = [
    ("미래에셋 거래내역증명서", detect_mirae_transactions),
    ("미래에셋 잔고증명서", detect_mirae_balance),
    ("토스증권 거래내역서", detect_toss_transactions),
    ("삼성증권 계좌거래내역 (주식보상)", detect_samsung_rsu),
    ("Robinhood Crypto statement", detect_robinhood_crypto_statement),
    ("Robinhood Gain/Loss report", detect_robinhood_gain_loss),
    ("consolidated 1099 (Fidelity / Chase / Robinhood)", detect_tax_form),
    ("빗썸 거래내역확인서", detect_bithumb_statement),
    ("빗썸 기간별 거래 내역 (.xlsx)", detect_bithumb_activity),
    ("Chase holdings / transactions CSV", detect_chase),
    ("Fidelity transactions CSV", detect_fidelity),
    ("Fidelity positions CSV", detect_fidelity_positions),
    ("Merrill holdings / transactions CSV", detect_merrill),
    ("Robinhood transactions CSV", detect_robinhood_transactions),
]


def identify(doc):
    """(label, Plan|Refusal) or (None, None).

    Every detector is asked, not just up to the first hit: two claims on one
    file means a marker is less specific than it looks, and that has to be an
    error rather than a coin toss decided by list order.
    """
    claims = []
    for label, detector in DETECTORS:
        verdict = detector(doc)
        if verdict is not None:
            claims.append((label, verdict))
    if not claims:
        return None, None
    if len(claims) > 1:
        names = ", ".join(label for label, _ in claims)
        raise RuntimeError(f"{doc.path.name}: claimed by more than one detector ({names})")
    return claims[0]


def number_split_families(plans):
    """Turn the members of one split statement into `-1of3`, `-2of3`, `-3of3`.

    Toss hands a long statement over as several PDFs that share a 발급번호 and a
    조회 기간; only their rows differ. The part numbers are therefore a property
    of the SET, and are worked out here once the run has seen all of it, in
    first-transaction order.

    A part dropped in on its own cannot be numbered — the file does not say it
    is 2 of 2 — so it keeps the plain name and, if a numbered set is already
    filed under that period, lands as a conflict for a person to look at rather
    than as a plausible-looking single file that silently replaces a set.
    """
    groups = {}
    for plan in plans.values():
        if plan.group:
            groups.setdefault(plan.group, []).append(plan)
    for members in groups.values():
        if len(members) < 2:
            continue
        members.sort(key=lambda p: p.order)
        total = len(members)
        for index, plan in enumerate(members, start=1):
            stem, dot, ext = plan.name.rpartition(".")
            plan.name = f"{stem}-{index}of{total}{dot}{ext}"
            plan.evidence.append(f"part {index} of {total} by first transaction date")


def existing_digests(directory):
    """{sha256: name} for everything already filed in a destination directory."""
    found = {}
    if not directory.is_dir():
        return found
    for path in sorted(directory.iterdir()):
        if path.is_file() and path.name not in IGNORED_NAMES:
            found[sha256_of(path)] = path.name
    return found


def ensure_inbox():
    INBOX_DIR.mkdir(parents=True, exist_ok=True)
    readme = INBOX_DIR / "README.md"
    current = readme.read_text(encoding="utf-8") if readme.exists() else None
    if current != INBOX_README:
        readme.write_text(INBOX_README, encoding="utf-8")


def inbox_files():
    return sorted(
        path
        for path in INBOX_DIR.iterdir()
        if path.is_file()
        and path.name not in IGNORED_NAMES
        and not path.name.startswith(".")
        and path.suffix.lower() not in IGNORED_SUFFIXES
    )


def main(argv):
    dry_run = "--dry-run" in argv or "-n" in argv
    unknown = [a for a in argv if a not in ("--dry-run", "-n", "-h", "--help")]
    if "-h" in argv or "--help" in argv or unknown:
        print(__doc__)
        print("usage: file-downloads.py [--dry-run]")
        return 2 if unknown else 0

    if not DATA_DIR.is_dir():
        print(f"ERROR: no data directory at {DATA_DIR}", file=sys.stderr)
        return 1
    ensure_inbox()

    files = inbox_files()
    header = f"inbox {INBOX_DIR}" + ("  (dry run)" if dry_run else "")
    print(header)
    if not files:
        print("  empty — nothing to file")
        return 0

    docs = {path: Document(path) for path in files}
    plans, refusals, unidentified = {}, [], []
    for path, doc in docs.items():
        _label, verdict = identify(doc)
        if verdict is None:
            unidentified.append(doc)
        elif isinstance(verdict, Refusal):
            refusals.append((doc, verdict))
        else:
            plans[path] = verdict
    number_split_families(plans)

    # Two inbox files planning the same destination are a collision this run
    # would resolve by whichever moved last. Neither moves.
    wanted = {}
    for path, plan in plans.items():
        wanted.setdefault((plan.subdir, plan.name), []).append(path)

    digests = {subdir: existing_digests(DATA_DIR / subdir) for subdir in {p.subdir for p in plans.values()}}

    filed, skipped, conflicts = [], [], []
    for path, plan in plans.items():
        destination = DATA_DIR / plan.subdir / plan.name
        shown = f"{plan.subdir}/{plan.name}"
        siblings = wanted[(plan.subdir, plan.name)]
        if len(siblings) > 1:
            others = ", ".join(sorted(p.name for p in siblings if p != path))
            conflicts.append((path, shown, plan, f"also claimed this run by {others}"))
            continue
        # Identity by CONTENT against the whole destination directory, not just
        # the target name: a document already filed under a different name is
        # still filed, and a second copy would be counted twice by the
        # extractors, which read every file they find.
        already = digests[plan.subdir].get(docs[path].digest)
        if already:
            skipped.append((path, f"{plan.subdir}/{already}", plan))
            continue
        if destination.exists():
            conflicts.append((path, shown, plan, "a different file is already filed under that name"))
            continue
        filed.append((path, destination, shown, plan))

    if filed:
        print("\nfiled")
        for path, destination, shown, plan in filed:
            print(f"  {path.name} → {shown}")
            for note in plan.evidence:
                print(f"      {note}")
            if docs[path].encrypted:
                print("      opened with STOCK_PDF_PASSWORD")
            if not dry_run:
                destination.parent.mkdir(parents=True, exist_ok=True)
                shutil.move(str(path), str(destination))

    if skipped:
        print("\nalready filed — left alone")
        for path, shown, _plan in skipped:
            print(f"  {path.name} — identical to {shown}")

    if conflicts:
        print("\nCONFLICT — not moved, resolve by hand")
        for path, shown, plan, why in conflicts:
            print(f"  {path.name} → {shown}: {why}")
            for note in plan.evidence:
                print(f"      {note}")

    if refusals:
        print("\nrecognised but not filed — left in the inbox")
        for doc, refusal in refusals:
            print(f"  {doc.path.name} — {refusal.what}: {refusal.reason}")
            if refusal.remedy:
                print(f"      {refusal.remedy}")

    if unidentified:
        print("\nunidentified — left in the inbox")
        for doc in unidentified:
            print(f"  {doc.path.name}")
            print(f"      first line: {doc.first_line()}")
            if doc.suffix == ".pdf" and doc.pdf_error:
                print(f"      could not be read as a PDF: {doc.pdf_error}")
            print(f"      no detector claimed it; looked for: "
                  f"{', '.join(label for label, _ in DETECTORS)}")

    remaining = len(skipped) + len(conflicts) + len(refusals) + len(unidentified)
    print(
        f"\n{len(filed)} filed{' (dry run — nothing moved)' if dry_run and filed else ''}, "
        f"{len(skipped)} already filed, {len(conflicts)} conflict(s), "
        f"{len(refusals) + len(unidentified)} unidentified — "
        f"{remaining + (len(filed) if dry_run else 0)} file(s) still in the inbox"
    )
    # Always 0 unless the run itself failed. An unidentified file is a message
    # to a person, not a broken pipeline, and it must not stop the hourly push
    # from carrying the files that WERE filed.
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
